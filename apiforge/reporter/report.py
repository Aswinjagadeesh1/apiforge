"""Report generation — Excel (pentest deliverable format) and JSON."""
from __future__ import annotations

import json
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apiforge.models import Finding

_SEVERITY_FILL = {
    "CRITICAL": "8B0000",
    "HIGH": "E63900",
    "MEDIUM": "E69500",
    "LOW": "E6C200",
    "INFO": "4A90A4",
}

_SEVERITY_ORDER = ["CRITICAL", "HIGH", "MEDIUM", "LOW", "INFO"]

_SUMMARY_COLUMNS = [
    ("S.No", 6),
    ("Severity", 12),
    ("CVSS", 8),
    ("Title", 38),
    ("OWASP Category", 38),
    ("CWE", 10),
    ("# EP", 6),
    ("Affected Endpoints", 46),
    ("Description", 55),
    ("Remediation", 50),
    ("PoC", 8),
]


class Reporter:
    def to_excel(
        self,
        findings: list[Finding],
        output_path: str | Path,
        target: str | None = None,
    ) -> None:
        import os
        import tempfile

        wb = Workbook()

        # ---- Sheet 1: Executive Summary (unchanged) ----
        summary = wb.active
        summary.title = "Executive Summary"
        self._build_summary(summary, findings, target)

        groups = group_findings(findings)

        # ---- Sheet 3: PoC detail (build first so we know anchor rows) ----
        poc_ws = wb.create_sheet("PoC")
        tmp = tempfile.mkdtemp(prefix="apiforge_xlsx_poc_")
        try:
            anchors = self._build_poc_sheet(poc_ws, groups, tmp)

            # ---- Sheet 2: Findings summary with POC hyperlinks ----
            findings_ws = wb.create_sheet("Findings")
            wb.move_sheet("Findings", -(len(wb.sheetnames) - 2))
            self._build_findings(findings_ws, groups, anchors)

            wb.save(str(output_path))
        finally:
            for fn in os.listdir(tmp):
                os.remove(os.path.join(tmp, fn))
            os.rmdir(tmp)

    def _build_summary(self, ws, findings: list[Finding], target: str | None) -> None:
        # Title banner
        ws.merge_cells("A1:D1")
        title = ws.cell(row=1, column=1, value="APIForge — API Security Assessment")
        title.font = Font(bold=True, size=16, color="FFFFFF")
        title.alignment = Alignment(horizontal="center", vertical="center")
        title.fill = PatternFill("solid", fgColor="1F3A5F")
        ws.row_dimensions[1].height = 30

        # Metadata block
        meta = [
            ("Target", target or "N/A"),
            ("Scan Date", datetime.now().strftime("%Y-%m-%d %H:%M")),
            ("Total Findings", str(len(group_findings(findings)))),
            ("Standard", "OWASP API Security Top 10 (2023)"),
        ]
        row = 3
        for label, value in meta:
            lc = ws.cell(row=row, column=1, value=label)
            lc.font = Font(bold=True)
            ws.cell(row=row, column=2, value=value)
            row += 1

        # Severity breakdown
        row += 1
        hdr = ws.cell(row=row, column=1, value="Findings by Severity")
        hdr.font = Font(bold=True, size=12, color="FFFFFF")
        hdr.fill = PatternFill("solid", fgColor="1F3A5F")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        row += 1

        sev_counts = Counter(g["severity"].value for g in group_findings(findings))
        for sev in _SEVERITY_ORDER:
            count = sev_counts.get(sev, 0)
            if count == 0:
                continue
            sc = ws.cell(row=row, column=1, value=sev)
            sc.fill = PatternFill("solid", fgColor=_SEVERITY_FILL.get(sev, "FFFFFF"))
            sc.font = Font(bold=True, color="FFFFFF")
            sc.alignment = Alignment(horizontal="center")
            ws.cell(row=row, column=2, value=count).alignment = Alignment(
                horizontal="center"
            )
            row += 1

        # OWASP category breakdown
        row += 1
        hdr2 = ws.cell(row=row, column=1, value="Findings by OWASP Category")
        hdr2.font = Font(bold=True, size=12, color="FFFFFF")
        hdr2.fill = PatternFill("solid", fgColor="1F3A5F")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        row += 1

        cat_counts = Counter(g["owasp"] for g in group_findings(findings))
        for cat, count in sorted(cat_counts.items()):
            ws.cell(row=row, column=1, value=cat).alignment = Alignment(wrap_text=True)
            ws.cell(row=row, column=2, value=count).alignment = Alignment(
                horizontal="center"
            )
            row += 1

        # Column widths
        ws.column_dimensions["A"].width = 42
        ws.column_dimensions["B"].width = 22
        ws.column_dimensions["C"].width = 12
        ws.column_dimensions["D"].width = 12

    def _build_findings(self, ws, groups, anchors) -> None:
        header_fill = PatternFill("solid", fgColor="1F3A5F")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        for col, (name, width) in enumerate(_SUMMARY_COLUMNS, 1):
            cell = ws.cell(row=1, column=col, value=name)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(col)].width = width

        for row, (i, g) in enumerate(enumerate(groups, 1), 2):
            values = [
                i, g["severity"].value, g["cvss"], g["title"], g["owasp"], g["cwe"],
                len(g["endpoints"]), "\n".join(g["endpoints"]),
                g["description"], g["remediation"],
            ]
            for col, value in enumerate(values, 1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            sev_cell = ws.cell(row=row, column=2)
            sev_cell.fill = PatternFill("solid", fgColor=_SEVERITY_FILL.get(g["severity"].value, "FFFFFF"))
            sev_cell.font = Font(bold=True, color="FFFFFF")
            poc = ws.cell(row=row, column=11, value="POC")
            poc.hyperlink = f"#PoC!A{anchors[i]}"
            poc.font = Font(bold=True, color="1155CC", underline="single")
            poc.alignment = Alignment(horizontal="center", vertical="center")

        ws.freeze_panes = "A2"

    def _build_poc_sheet(self, ws, groups, tmp) -> dict:
        """PoC detail sheet: per group, the vuln name + one annotated screenshot
        per affected endpoint. Returns {i: anchor_row} for the summary links."""
        import os

        from openpyxl.drawing.image import Image as XLImage

        from apiforge.reporter.poc_render import render_poc

        anchors: dict = {}
        prow = 1
        for i, g in enumerate(groups, 1):
            anchors[i] = prow
            title = ws.cell(row=prow, column=1, value=f"{i}.  {g['title']}")
            title.font = Font(bold=True, size=13, color="1F3A5F")
            prow += 1
            for j, m in enumerate(g["members"], 1):
                lbl = ws.cell(row=prow, column=1, value=f"{m.method} {m.endpoint}")
                lbl.font = Font(italic=True, size=10, color="555555")
                prow += 1
                img_path = os.path.join(tmp, f"poc_{i}_{j}.png")
                render_poc(
                    m.poc_request, m.poc_response, img_path,
                    highlight=getattr(m, "evidence", None),
                    caption=getattr(m, "evidence_caption", None),
                )
                img = XLImage(img_path)
                img.width = int(img.width * 0.62)
                img.height = int(img.height * 0.62)
                ws.add_image(img, f"A{prow}")
                prow += max(18, int(img.height / 20) + 2) + 1
            prow += 2
        return anchors

    def to_json(self, findings: list[Finding], output_path: str | Path) -> None:
        data = [f.model_dump() for f in findings]
        with open(output_path, "w", encoding="utf-8") as fh:
            json.dump(data, fh, indent=2, default=str)


# ---- Word (.docx) export -------------------------------------------------
# Added as a separate function to keep the Reporter class edit minimal.

def write_word_report(findings, output_path, target=None):
    """Generate a professional .docx pentest report."""
    from collections import Counter
    from datetime import datetime

    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    _SEV_RGB = {
        "CRITICAL": RGBColor(0x8B, 0x00, 0x00),
        "HIGH": RGBColor(0xE6, 0x39, 0x00),
        "MEDIUM": RGBColor(0xE6, 0x95, 0x00),
        "LOW": RGBColor(0xE6, 0xC2, 0x00),
        "INFO": RGBColor(0x4A, 0x90, 0xA4),
    }
    _ORDER = ["CRITICAL", "HIGH", "MEDIUM", "LOW", "INFO"]

    doc = Document()

    # ---- Title page ----
    title = doc.add_heading("API Security Assessment", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("Generated by APIForge — OWASP API Top 10 Scanner")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    meta = doc.add_paragraph()
    meta.add_run(f"Target: ").bold = True
    meta.add_run(f"{target or 'N/A'}\n")
    meta.add_run(f"Scan Date: ").bold = True
    meta.add_run(f"{datetime.now().strftime('%Y-%m-%d %H:%M')}\n")
    meta.add_run(f"Total Findings: ").bold = True
    meta.add_run(f"{len(findings)}\n")
    meta.add_run(f"Standard: ").bold = True
    meta.add_run("OWASP API Security Top 10 (2023)")

    # ---- Executive summary ----
    doc.add_heading("Executive Summary", level=1)

    sev_counts = Counter(f.severity.value for f in findings)
    doc.add_heading("Findings by Severity", level=2)
    t = doc.add_table(rows=1, cols=2)
    t.style = "Light Grid Accent 1"
    t.rows[0].cells[0].text = "Severity"
    t.rows[0].cells[1].text = "Count"
    for sev in _ORDER:
        c = sev_counts.get(sev, 0)
        if c == 0:
            continue
        row = t.add_row().cells
        run = row[0].paragraphs[0].add_run(sev)
        run.bold = True
        run.font.color.rgb = _SEV_RGB.get(sev, RGBColor(0, 0, 0))
        row[1].text = str(c)

    cat_counts = Counter(f.owasp_category for f in findings)
    doc.add_heading("Findings by OWASP Category", level=2)
    t2 = doc.add_table(rows=1, cols=2)
    t2.style = "Light Grid Accent 1"
    t2.rows[0].cells[0].text = "OWASP Category"
    t2.rows[0].cells[1].text = "Count"
    for cat, c in sorted(cat_counts.items()):
        row = t2.add_row().cells
        row[0].text = cat
        row[1].text = str(c)

    doc.add_page_break()

    # ---- Detailed findings ----
    doc.add_heading("Detailed Findings", level=1)
    ordered = sorted(findings, key=lambda f: (f.severity.rank, f.check_id))
    for i, f in enumerate(ordered, 1):
        h = doc.add_heading(f"{i}. {f.title}", level=2)

        sev_p = doc.add_paragraph()
        sev_run = sev_p.add_run(f"{f.severity.value}")
        sev_run.bold = True
        sev_run.font.color.rgb = _SEV_RGB.get(f.severity.value, RGBColor(0, 0, 0))
        sev_p.add_run(f"  |  CVSS {f.cvss_score}  |  {f.cwe}  |  {f.owasp_category}")

        meta_p = doc.add_paragraph()
        meta_p.add_run("Endpoint: ").bold = True
        meta_p.add_run(f"{f.method} {f.endpoint}")

        doc.add_paragraph().add_run("Description").bold = True
        doc.add_paragraph(f.description)

        doc.add_paragraph().add_run("Proof of Concept — Request").bold = True
        pr = doc.add_paragraph(f.poc_request)
        pr.style = "No Spacing"

        doc.add_paragraph().add_run("Proof of Concept — Response").bold = True
        rr = doc.add_paragraph(f.poc_response)
        rr.style = "No Spacing"

        doc.add_paragraph().add_run("Remediation").bold = True
        doc.add_paragraph(f.remediation)

        doc.add_paragraph("─" * 40)

    doc.save(str(output_path))


def group_findings(findings):
    """Group findings by check so one root-cause vuln across many endpoints is
    reported once, with all affected endpoints, not N separate findings."""
    from collections import OrderedDict
    buckets = OrderedDict()
    for f in findings:
        buckets.setdefault(f.check_id, []).append(f)
    groups = []
    for members in buckets.values():
        rep = min(members, key=lambda m: m.severity.rank)
        seen, eps = set(), []
        for m in members:
            key = f"{m.method} {m.endpoint}"
            if key not in seen:
                seen.add(key)
                eps.append(key)
        if len(members) > 1:
            desc = (f"This vulnerability was identified on {len(eps)} endpoints "
                    f"(listed under Affected Endpoints). " + rep.description)
        else:
            desc = rep.description
        groups.append({
            "title": rep.title,
            "severity": rep.severity,
            "cvss": max(m.cvss_score for m in members),
            "owasp": rep.owasp_category,
            "cwe": rep.cwe,
            "endpoints": eps,
            "description": desc,
            "remediation": rep.remediation,
            "members": members,
        })
    groups.sort(key=lambda g: g["severity"].rank)
    return groups
