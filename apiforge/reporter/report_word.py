"""apiforge/reporter/report_word.py

Adiroha-format Word (.docx) report generator for APIForge.

Drop-in replacement for the old write_word_report() in report.py — same
name and signature (findings, output_path, target=None), plus an optional
`meta` dict for per-engagement cover details, so cli.py needs no change
beyond pointing its import here.

Consumes the real Finding model directly:
  check_id, title, severity (Enum), owasp_category, cwe, cvss_score,
  method, endpoint, description, poc_request, poc_response, remediation
The red-box highlight uses finding.evidence if a check sets it, otherwise a
heuristic auto-highlight. No status field is assumed (assessment report);
set meta["is_reassessment"]=True only if your findings carry a .status.
"""
from __future__ import annotations

import os
import tempfile
from collections import Counter
from datetime import date

from apiforge.reporter.poc_render import render_poc

_HDR_BLUE = "1F4E79"
_SEV_HEX = {
    "CRITICAL": "C00000", "HIGH": "D62828", "MEDIUM": "E08E0B",
    "LOW": "2E7D32", "INFO": "4A90A4",
}
_SEV_ORDER = ["CRITICAL", "HIGH", "MEDIUM", "LOW", "INFO"]

_OWASP_API_TOP10 = [
    ("API1:2023", "Broken Object Level Authorization"),
    ("API2:2023", "Broken Authentication"),
    ("API3:2023", "Broken Object Property Level Authorization"),
    ("API4:2023", "Unrestricted Resource Consumption"),
    ("API5:2023", "Broken Function Level Authorization"),
    ("API6:2023", "Unrestricted Access to Sensitive Business Flows"),
    ("API7:2023", "Server Side Request Forgery"),
    ("API8:2023", "Security Misconfiguration"),
    ("API9:2023", "Improper Inventory Management"),
    ("API10:2023", "Unsafe Consumption of APIs"),
]

_DEFAULT_META = {
    "report_title": "API Security Assessment Report",
    "is_reassessment": False,
    "vendor_name": "Company Name",
    "vendor_address": "Company Address",
    "vendor_contact": "",
    "test_type": "Grey Box Testing",
    "classification": "Client Confidential",
    "client_name": "",
    "scope_items": None,          # defaults to [target]
    "assessment_period": date.today().strftime("%d-%m-%Y"),
}


def write_word_report(findings, output_path, target=None, meta=None):
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    M = dict(_DEFAULT_META)
    if meta:
        M.update(meta)
    if not M.get("scope_items"):
        M["scope_items"] = [target or "In-scope APIs"]

    tmp = tempfile.mkdtemp(prefix="apiforge_poc_")
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    # ---- small helpers ----
    def center(text, size, color=None, bold=True, italic=False):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)

    def h1(text):
        r = doc.add_heading(level=1).add_run(text)
        r.font.color.rgb = RGBColor.from_string(_HDR_BLUE)

    def h2(text):
        r = doc.add_heading(level=2).add_run(text)
        r.font.color.rgb = RGBColor.from_string(_HDR_BLUE)

    def field_label(text):
        p = doc.add_paragraph()
        r = p.add_run(text.upper()); r.bold = True; r.font.size = Pt(11)
        r.font.color.rgb = RGBColor.from_string(_HDR_BLUE)
        pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
        b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "6"); b.set(qn("w:space"), "1")
        b.set(qn("w:color"), _HDR_BLUE); pbdr.append(b); pPr.append(pbdr)

    def value(text, color=None, bold=False, mono=False):
        p = doc.add_paragraph()
        if mono:
            p.style = "No Spacing"
        r = p.add_run(text); r.bold = bold
        if mono:
            r.font.name = "Consolas"; r.font.size = Pt(8.5)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)

    def toc():
        p = doc.add_paragraph(); run = p.add_run()
        parts = [("fld", "begin"), ("instr", 'TOC \\o "1-2" \\h \\z \\u'),
                 ("fld", "separate"), ("txt", "Update this field in Word (F9)."), ("fld", "end")]
        for kind, val in parts:
            if kind == "instr":
                e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = val
            elif kind == "txt":
                e = OxmlElement("w:t"); e.text = val
            else:
                e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), val)
            run._r.append(e)

    ordered = sorted(findings, key=lambda f: (f.severity.rank, f.check_id))
    counts = Counter(f.severity.value for f in findings)

    # ---- COVER ----
    doc.add_paragraph("\n\n")
    center(M["report_title"], 26, _HDR_BLUE)
    center(target or M.get("client_name") or "", 15, bold=False)
    doc.add_paragraph("\n")
    center("Prepared by " + M["vendor_name"], 12, bold=True)
    center(M["vendor_address"], 10, bold=False, italic=True)
    center(M["vendor_contact"], 10, bold=False)
    doc.add_paragraph("\n")
    center(f'Classification: {M["classification"]}   ·   {M["test_type"]}', 10, italic=True, bold=False)
    center("Report auto-generated by APIForge", 9, italic=True, bold=False, color="888888")
    doc.add_page_break()

    # ---- TOC ----
    h1("Table of Contents"); toc(); doc.add_page_break()

    # ---- 1 INTRODUCTION ----
    h1("1. Introduction")
    h2("1.1 Summary")
    kind = "re-assessment" if M["is_reassessment"] else "assessment"
    value(f'This report presents the vulnerability {kind} for {target or "the in-scope APIs"}, '
          f'conducted as {M["test_type"]} against the OWASP API Security Top 10 (2023). Testing '
          f'assumed the perspective of a malicious user while taking due care not to harm the application.')
    h2("1.2 Aim")
    value(f'To assess the security posture of {target or "the in-scope APIs"} and provide actionable '
          f'remediation for each confirmed finding.')
    h2("1.3 Assessment Objective")
    value("Identify, verify and rate security flaws in the in-scope APIs against industry-standard benchmarks.")

    # ---- 2 SCOPE ----
    h1("2. Scope of Assessment")
    h2("2.1 Scope of Testing")
    t = doc.add_table(rows=1, cols=2); t.style = "Light Grid Accent 1"
    t.rows[0].cells[0].text = "S.No"; t.rows[0].cells[1].text = "Scope of Testing"
    for c in t.rows[0].cells:
        c.paragraphs[0].runs[0].bold = True
    for i, item in enumerate(M["scope_items"], 1):
        r = t.add_row().cells; r[0].text = str(i); r[1].text = str(item)

    # ---- 3 TERMS & LEGEND ----
    h1("3. Terms, Definitions and Legends")
    h2("3.1 Vulnerability Report Format")
    value("Each finding in Section 5 is reported with: Severity, Affected Endpoint, "
          + ("Status After Reassessment, " if M["is_reassessment"] else "")
          + "OWASP Category, CVSS, Description, Proof of Concept (annotated evidence), and Remediation.")
    h2("3.2 Severity Levels")
    t = doc.add_table(rows=1, cols=2); t.style = "Light Grid Accent 1"
    for c, hh in zip(t.rows[0].cells, ["Severity", "Meaning"]):
        c.text = hh; c.paragraphs[0].runs[0].bold = True
    for sev, meaning in [
        ("CRITICAL", "Severe impact, trivial to exploit — fix urgently."),
        ("HIGH", "Significant impact on confidentiality/integrity — prioritise."),
        ("MEDIUM", "Moderate impact or requires some preconditions."),
        ("LOW", "Limited impact — hardening / best-practice."),
        ("INFO", "Informational; no direct security impact."),
    ]:
        cells = t.add_row().cells
        cells[0].text = sev
        cells[0].paragraphs[0].runs[0].bold = True
        cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(_SEV_HEX[sev])
        cells[1].text = meaning

    # ---- 4 EXECUTIVE SUMMARY ----
    h1("4. Executive Summary of Vulnerabilities")
    h2("4.1 Vulnerabilities Ordered by Severity")
    t = doc.add_table(rows=1, cols=5); t.style = "Light Grid Accent 1"
    for c, hh in zip(t.rows[0].cells, ["S.No", "Vulnerability", "Severity", "CVSS", "OWASP"]):
        c.text = hh; c.paragraphs[0].runs[0].bold = True
    for i, f in enumerate(ordered, 1):
        cells = t.add_row().cells
        cells[0].text = str(i); cells[1].text = f.title
        cells[2].text = f.severity.value
        cells[2].paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(_SEV_HEX.get(f.severity.value, "000000"))
        cells[2].paragraphs[0].runs[0].bold = True
        cells[3].text = str(f.cvss_score)
        cells[4].text = (f.owasp_category or "").split(" ")[0].split(":")[0] or f.owasp_category
    h2("4.2 Total Vulnerabilities by Severity")
    present = [s for s in _SEV_ORDER if counts.get(s, 0)]
    t = doc.add_table(rows=1, cols=len(present) + 1); t.style = "Light Grid Accent 1"
    for c, hh in zip(t.rows[0].cells, ["Total"] + present):
        c.text = hh; c.paragraphs[0].runs[0].bold = True
    row = t.add_row().cells
    row[0].text = str(len(findings)); row[0].paragraphs[0].runs[0].bold = True
    for j, sev in enumerate(present, 1):
        row[j].text = str(counts[sev])
        r = row[j].paragraphs[0].runs[0]; r.bold = True
        r.font.color.rgb = RGBColor.from_string(_SEV_HEX[sev])
    doc.add_page_break()

    # ---- 5 FINDINGS ----
    h1("5. Vulnerabilities List with Remediation")
    for i, f in enumerate(ordered, 1):
        h2(f"5.{i}  {f.title}")
        field_label("Severity"); value(f.severity.value, _SEV_HEX.get(f.severity.value), bold=True)
        if M["is_reassessment"]:
            status = getattr(f, "status", "Open")
            field_label("Status After Reassessment")
            value(status, "2E7D32" if str(status).lower() == "closed" else "D62828", bold=True)
        field_label("Affected Endpoint"); value(f"{f.method} {f.endpoint}")
        field_label("OWASP Category"); value(f"{f.owasp_category}   (CVSS {f.cvss_score}, {f.cwe})")
        field_label("Description"); value(f.description)

        field_label("Proof of Concept")
        img_path = os.path.join(tmp, f"poc_{i}.png")
        render_poc(
            f.poc_request, f.poc_response, img_path,
            highlight=getattr(f, "evidence", None),
            caption=getattr(f, "evidence_caption", None),
        )
        pic = doc.add_paragraph(); pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.add_run().add_picture(img_path, width=Inches(6.3))
        cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = cap.add_run(f"Figure 5.{i}: annotated request/response evidence for {f.title}.")
        cr.italic = True; cr.font.size = Pt(9)

        field_label("Remediation"); value(f.remediation)
        doc.add_page_break()

    # ---- Tools / Conclusion / Way Ahead ----
    h1("Tools and References")
    for tool in ["APIForge (automated OWASP API pass)", "Postman", "Burp Suite", "jwt_tool", "curl"]:
        value("• " + tool)
    ch = counts.get("CRITICAL", 0) + counts.get("HIGH", 0)
    h1("Conclusion")
    value(f'The assessment identified {len(findings)} finding(s), {ch} of Critical/High severity. '
          f'The most serious issues concern authentication and object-level authorization — the flaw '
          f'classes behind most real-world API breaches. Prioritise remediation of Critical/High items.')
    h1("Way Ahead")
    value("Remediate in severity order, then request a re-assessment to confirm closure. "
          "Adopt the OWASP API Security Top 10 as a recurring control baseline.")

    # ---- Appendix ----
    h1("Appendix A: OWASP API Top 10 (2023) Coverage")
    seen = {(f.owasp_category or "").split(":")[0].split(" ")[0] for f in findings}
    t = doc.add_table(rows=1, cols=3); t.style = "Light Grid Accent 1"
    for c, hh in zip(t.rows[0].cells, ["Category", "Name", "Findings in this report"]):
        c.text = hh; c.paragraphs[0].runs[0].bold = True
    for cat, name in _OWASP_API_TOP10:
        cells = t.add_row().cells
        cells[0].text = cat; cells[1].text = name
        hit = cat.split(":")[0] in seen
        cells[2].text = "Yes" if hit else "—"
        if hit:
            cells[2].paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string("2E7D32")

    doc.save(str(output_path))
    # PoC PNGs are already embedded; clean the temp dir
    for fn in os.listdir(tmp):
        os.remove(os.path.join(tmp, fn))
    os.rmdir(tmp)
    return str(output_path)
